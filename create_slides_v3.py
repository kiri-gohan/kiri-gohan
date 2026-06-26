#!/usr/bin/env python3
"""
ファスティング指導士養成講座 カウンセリング講義スライド v3
カラー：温かみのあるベージュ系 / フォント：明朝体 / 絵文字なし
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ─── カラーパレット（ベージュ系）────────────────────────
C_BG      = RGBColor(0xFB, 0xF7, 0xF2)   # 温かみのあるオフホワイト
C_BG2     = RGBColor(0xF5, 0xEE, 0xE4)   # 薄いベージュ
C_SAND    = RGBColor(0xE8, 0xD8, 0xC4)   # サンドベージュ
C_WARM    = RGBColor(0xD6, 0xC2, 0xA8)   # ウォームベージュ（アクセント）
C_MOCHA   = RGBColor(0x8B, 0x6F, 0x55)   # モカブラウン（見出し）
C_BROWN   = RGBColor(0x5C, 0x3D, 0x2E)   # ダークブラウン（タイトル）
C_DGRAY   = RGBColor(0x3D, 0x35, 0x30)   # 濃い茶グレー（本文）
C_MGRAY   = RGBColor(0x7A, 0x6A, 0x5E)   # 中間茶グレー（補足）
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_GOLD    = RGBColor(0xB8, 0x8B, 0x4A)   # ゴールド（強調）
C_SAGE    = RGBColor(0x7A, 0x96, 0x74)   # セージグリーン
C_RUST    = RGBColor(0xA0, 0x5A, 0x4A)   # 控えめな赤錆（注意・対比）
C_LINE    = RGBColor(0xC8, 0xB8, 0xA4)   # ライン用ベージュ

MINCHO = "游明朝"  # Windows/Mac共通の明朝体
MINCHO_ALT = "ＭＳ 明朝"


def set_font_mincho(run):
    """ランに明朝体を設定する"""
    rPr = run._r.get_or_add_rPr()
    # 日本語フォント
    latin = etree.SubElement(rPr, qn('a:latin'))
    latin.set('typeface', MINCHO)
    ea = etree.SubElement(rPr, qn('a:ea'))
    ea.set('typeface', MINCHO)


def rect(slide, l, t, w, h, fill=None, line=None, lw=1.5):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    return s


def txt(slide, text, l, t, w, h, size=16, bold=False, color=None,
        align=PP_ALIGN.LEFT, italic=False, mincho=True):
    if color is None:
        color = C_DGRAY
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if mincho:
        set_font_mincho(r)
    return tb


def mtxt(slide, lines, l, t, w, h, size=15, bold=False, color=None,
         align=PP_ALIGN.LEFT, line_colors=None, mincho=True):
    if color is None:
        color = C_DGRAY
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = line_colors[i] if line_colors else color
        if mincho:
            set_font_mincho(r)
    return tb


def header(slide, title_text, bar_color=None):
    """統一ヘッダー帯"""
    if bar_color is None:
        bar_color = C_MOCHA
    rect(slide, 0, 0, 13.33, 7.5, fill=C_BG)
    rect(slide, 0, 0, 13.33, 1.05, fill=bar_color)
    # 細いゴールドライン
    rect(slide, 0, 1.05, 13.33, 0.06, fill=C_GOLD)
    txt(slide, title_text, 0.5, 0.15, 12.5, 0.82,
        size=26, bold=True, color=C_WHITE, align=PP_ALIGN.LEFT)


# ══════════════════════════════════════════════════════════
# スライド 1 ： タイトル
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.33, 7.5, fill=C_BG2)

# 上下の装飾帯
rect(s, 0, 0, 13.33, 0.25, fill=C_BROWN)
rect(s, 0, 7.25, 13.33, 0.25, fill=C_BROWN)
rect(s, 0, 0.25, 13.33, 0.06, fill=C_GOLD)
rect(s, 0, 7.19, 13.33, 0.06, fill=C_GOLD)

# 左右の縦帯
rect(s, 0, 0.25, 0.55, 6.94, fill=C_SAND)
rect(s, 0.55, 0.25, 0.06, 6.94, fill=C_WARM)
rect(s, 12.72, 0.25, 0.55, 6.94, fill=C_SAND)
rect(s, 12.66, 0.25, 0.06, 6.94, fill=C_WARM)

# タイトル枠
rect(s, 1.6, 1.5, 10.1, 2.8, fill=C_WHITE, line=C_WARM, lw=2)
rect(s, 1.6, 1.5, 10.1, 0.1, fill=C_MOCHA)
rect(s, 1.6, 4.2, 10.1, 0.1, fill=C_MOCHA)

txt(s, "カウンセリングの本質", 1.8, 1.7, 9.7, 1.3,
    size=46, bold=True, color=C_BROWN, align=PP_ALIGN.CENTER)

rect(s, 3.5, 3.0, 6.3, 0.05, fill=C_MOCHA)

txt(s, "── 答えを教えるのではなく、気づきを生む対話へ ──",
    1.8, 3.1, 9.7, 0.75, size=18, color=C_MOCHA,
    align=PP_ALIGN.CENTER, italic=True)

txt(s, "ファスティング指導士養成講座  ／  カウンセリング講義",
    2.5, 4.5, 8.3, 0.6, size=14, color=C_MGRAY, align=PP_ALIGN.CENTER)

rect(s, 2.0, 5.3, 9.3, 0.95, fill=C_SAND, line=C_WARM, lw=1.5)
txt(s, "「悩みに気づかせる場が、最大の教育になる」",
    2.2, 5.42, 8.9, 0.72, size=17, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 2 ： アジェンダ
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "本日の講義内容")

items = [
    ("01", "カウンセリングとは？", "よくある誤解を解く"),
    ("02", "カウンセリングの真の目的", "「気づき」を生む対話"),
    ("03", "悩みの構造を理解する", "なぜ一人では気づけないのか"),
    ("04", "なりたい自分を見つける", "目標設定のプロセス"),
    ("05", "Have to → Want to", "内発的な動機を引き出す"),
    ("06", "一緒に叶える環境と知識", "サポートの重要性"),
    ("07", "カウンセリングは最大の教育", "気づかせる力"),
    ("08", "実践：問いかけの技術", "聴く力・問う力"),
]

for i, (num, title, sub) in enumerate(items):
    row, col = i // 2, i % 2
    x, y = 0.4 + col * 6.45, 1.28 + row * 1.48
    rect(s, x, y, 6.1, 1.35, fill=C_WHITE, line=C_LINE, lw=1.5)
    rect(s, x, y, 0.85, 1.35, fill=C_SAND)
    txt(s, num, x + 0.05, y + 0.42, 0.75, 0.5, size=18, bold=True,
        color=C_MOCHA, align=PP_ALIGN.CENTER)
    rect(s, x + 0.85, y, 5.25, 0.58, fill=C_BG2)
    txt(s, title, x + 1.0, y + 0.08, 5.0, 0.45,
        size=15, bold=True, color=C_BROWN)
    txt(s, sub, x + 1.0, y + 0.72, 5.0, 0.55, size=13, color=C_MGRAY)


# ══════════════════════════════════════════════════════════
# スライド 3 ： よくある誤解 vs 本質
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "01｜カウンセリングとは？")

# 左：誤解
rect(s, 0.4, 1.2, 5.8, 5.75, fill=C_WHITE, line=RGBColor(0xC8, 0x9A, 0x8E), lw=1.5)
rect(s, 0.4, 1.2, 5.8, 0.65, fill=RGBColor(0xC8, 0x9A, 0x8E))
txt(s, "よくある思い込み", 0.55, 1.27, 5.5, 0.52,
    size=17, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

miss = [
    "「全部 答えてあげなきゃ」",
    "「全部 伝えてあげなきゃ」",
    "「知識を教えることが\nカウンセリング」",
    "「専門家が正解を提供する場」",
]
for j, m in enumerate(miss):
    y = 2.0 + j * 1.1
    rect(s, 0.6, y, 5.4, 0.98, fill=RGBColor(0xFD, 0xF2, 0xF0),
         line=RGBColor(0xC8, 0x9A, 0x8E), lw=1)
    txt(s, "×", 0.65, y + 0.2, 0.5, 0.6, size=20, bold=True,
        color=C_RUST, align=PP_ALIGN.CENTER)
    txt(s, m, 1.2, y + 0.15, 4.6, 0.72, size=14, bold=True, color=C_RUST)

txt(s, "→", 6.35, 3.85, 0.7, 0.7, size=36, bold=True,
    color=C_MOCHA, align=PP_ALIGN.CENTER)

# 右：本質
rect(s, 7.1, 1.2, 5.8, 5.75, fill=C_WHITE, line=C_WARM, lw=1.5)
rect(s, 7.1, 1.2, 5.8, 0.65, fill=C_MOCHA)
txt(s, "カウンセリングの本質", 7.25, 1.27, 5.5, 0.52,
    size=17, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

truths = [
    "「気づきを生む場」",
    "「なりたい自分を見つける場」",
    "「悩みの本質に\n自分で気づく場」",
    "「最大の教育ができる場」",
]
for j, t in enumerate(truths):
    y = 2.0 + j * 1.1
    rect(s, 7.3, y, 5.4, 0.98, fill=C_BG2, line=C_LINE, lw=1)
    txt(s, "◆", 7.35, y + 0.2, 0.55, 0.6, size=16, bold=True,
        color=C_GOLD, align=PP_ALIGN.CENTER)
    txt(s, t, 7.95, y + 0.15, 4.65, 0.72, size=14, bold=True, color=C_BROWN)


# ══════════════════════════════════════════════════════════
# スライド 4 ： カウンセリングの真の目的
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "02｜カウンセリングの真の目的")

rect(s, 0.8, 1.2, 11.7, 0.95, fill=C_SAND, line=C_WARM, lw=1.5)
txt(s, "「答えを教える場」ではなく\n「自分に足りないものに気づかせる場」",
    1.0, 1.25, 11.3, 0.9, size=20, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)

goals = [
    ("気づきを生む",
     "クライアントが自分では見えていない\n悩みの本質・原因に気づく"),
    ("なりたい自分を\n見つける",
     "理想の未来像を明確にして\nそこへ向かう意欲を引き出す"),
    ("行動への\n橋渡し",
     "気づきを一人で抱えず\n必要な環境・知識につなげる"),
]

for i, (title, desc) in enumerate(goals):
    x = 0.6 + i * 4.2
    rect(s, x, 2.4, 3.9, 4.65, fill=C_WHITE, line=C_LINE, lw=1.5)
    rect(s, x, 2.4, 3.9, 0.65, fill=C_SAND)
    rect(s, x, 3.05, 3.9, 0.62, fill=C_MOCHA)
    txt(s, title, x + 0.1, 3.09, 3.7, 0.58,
        size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    txt(s, str(i + 1), x + 1.5, 2.45, 0.9, 0.6, size=24, bold=True,
        color=C_MOCHA, align=PP_ALIGN.CENTER)
    txt(s, desc, x + 0.15, 3.82, 3.6, 3.1,
        size=14, color=C_DGRAY, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 5 ： なぜ一人では気づけないのか
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "03｜悩みの構造 ── なぜ一人では気づけないの？",
       bar_color=RGBColor(0x6E, 0x58, 0x48))

# 左：氷山モデル
rect(s, 0.4, 1.2, 5.6, 5.85, fill=C_WHITE, line=C_LINE, lw=1.5)
txt(s, "氷山モデル：悩みの構造", 0.6, 1.28, 5.2, 0.5,
    size=14, bold=True, color=C_MOCHA, align=PP_ALIGN.CENTER)

layers = [
    (RGBColor(0xEE, 0xD8, 0xC8), "水面上（見えている悩み）",
     "「体重が減らない」「続かない」", C_BROWN),
    (RGBColor(0xD8, 0xBE, 0xA8), "水面下（隠れた本音）",
     "「自己否定感」「孤独感」\n「誰にも言えないストレス」", C_BROWN),
    (RGBColor(0x8B, 0x6F, 0x55), "深部（根本原因）",
     "生活習慣の背景にある\nストレス・環境・\nセルフイメージの問題", C_WHITE),
]
layer_h = [0.98, 1.25, 1.78]
y_pos = 1.85
for j, ((bg, label, desc, tc), lh) in enumerate(zip(layers, layer_h)):
    rect(s, 0.6, y_pos, 5.2, lh, fill=bg)
    txt(s, label, 0.75, y_pos + 0.06, 5.0, 0.42,
        size=13, bold=True, color=tc)
    txt(s, desc, 0.75, y_pos + 0.5, 5.0, lh - 0.55,
        size=12, color=tc, align=PP_ALIGN.CENTER)
    y_pos += lh + 0.04

# 右：3つの理由
txt(s, "一人では気づけない　3つの理由", 6.3, 1.25, 6.65, 0.55,
    size=17, bold=True, color=C_BROWN)
rect(s, 6.3, 1.83, 6.7, 0.05, fill=C_GOLD)

reasons = [
    ("客観視が難しい",
     "自分の「当たり前」が問題だと\n気づきにくい。魚は水が見えない。"),
    ("思い込みのフィルター",
     "「どうせ無理」「自分はダメ」という\n固定観念が本質を覆い隠す。"),
    ("同じ視点での堂々巡り",
     "同じ思考パターンで考え続けても\n新しい気づきは生まれにくい。"),
]
for j, (title, desc) in enumerate(reasons):
    y = 2.03 + j * 1.72
    rect(s, 6.3, y, 6.7, 1.6, fill=C_WHITE, line=C_LINE, lw=1)
    rect(s, 6.3, y, 6.7, 0.55, fill=C_BG2)
    txt(s, f"{j+1}", 6.35, y + 0.1, 0.55, 0.4, size=18, bold=True,
        color=C_GOLD, align=PP_ALIGN.CENTER)
    txt(s, title, 6.95, y + 0.08, 5.9, 0.42, size=14, bold=True, color=C_BROWN)
    txt(s, desc, 6.45, y + 0.63, 6.4, 0.92, size=13, color=C_DGRAY)


# ══════════════════════════════════════════════════════════
# スライド 6 ： 気づきが生まれる対話例
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "「気づき」が生まれる瞬間 ── 対話のリアル",
       bar_color=RGBColor(0x6E, 0x58, 0x48))

txt(s, "カウンセラーの問いかけが、クライアントの内側にある答えを引き出す",
    0.5, 1.15, 12.3, 0.55, size=15, bold=True,
    color=C_MOCHA, align=PP_ALIGN.CENTER, italic=True)

dialogs = [
    ("カウンセラー",
     "「ファスティングで一番うまくいかないのは、どんなときですか？」",
     C_MOCHA, C_BG2),
    ("クライアント",
     "「夜、仕事から帰ると… つい食べてしまうんです」",
     C_MGRAY, C_WHITE),
    ("カウンセラー",
     "「そのとき、どんな気持ちになっていますか？」",
     C_MOCHA, C_BG2),
    ("クライアント",
     "「疲れていて… 自分へのご褒美が欲しくなる感じがします」",
     C_MGRAY, C_WHITE),
    ("カウンセラー",
     "「食べること以外に、何かご褒美があったら変わりそうですか？」",
     C_MOCHA, C_BG2),
    ("クライアント",
     "「……あ！ そっか。食べること以外で自分を癒す方法、考えたことなかったです」",
     C_BROWN, C_SAND),
]

for j, (who, line, col, bg) in enumerate(dialogs):
    y = 1.82 + j * 0.86
    rect(s, 0.4, y, 12.5, 0.82, fill=bg, line=C_LINE, lw=0.5)
    is_c = who == "カウンセラー"
    lc = C_MOCHA if is_c else C_MGRAY
    txt(s, f"【{who}】", 0.52, y + 0.16, 2.0, 0.5, size=11, bold=True, color=lc)
    bold = (j == 5)
    txt(s, line, 2.55, y + 0.12, 10.1, 0.66, size=14, bold=bold, color=col)

rect(s, 0.4, 7.1, 12.5, 0.32, fill=C_SAND)
txt(s, "気づきは「教えられるもの」ではなく、問いかけによって自分の中から生まれるもの",
    0.6, 7.12, 12.0, 0.28, size=13, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 7 ： なりたい自分を見つける 4ステップ
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "04｜なりたい自分を見つける ── 目標設定の4ステップ")

steps = [
    ("STEP 1", "現在地を知る",
     "今、どんな状態？\n何に困っている？\n何を感じている？",
     RGBColor(0xEA, 0xD5, 0xC4)),
    ("STEP 2", "痛みを掘り下げる",
     "その悩みの奥にある\n本当の感情は？\nいつからそう感じてる？",
     RGBColor(0xDC, 0xCB, 0xBC)),
    ("STEP 3", "理想を描く",
     "どんな自分になりたい？\nそれが叶ったらどんな気持ち？\n大切にしたい価値観は？",
     RGBColor(0xCC, 0xD8, 0xC0)),
    ("STEP 4", "ギャップを言語化",
     "現在地と理想の差は？\n何があれば埋まる？\n何が邪魔している？",
     C_SAND),
]

for i, (step, title, desc, bg) in enumerate(steps):
    x = 0.45 + i * 3.2
    rect(s, x, 1.25, 3.05, 5.75, fill=C_WHITE, line=C_LINE, lw=1.5)
    rect(s, x, 1.25, 3.05, 0.7, fill=bg)
    txt(s, step, x + 0.05, 1.28, 2.95, 0.4, size=12, bold=True,
        color=C_BROWN, align=PP_ALIGN.CENTER)
    rect(s, x, 1.95, 3.05, 0.62, fill=C_MOCHA)
    txt(s, title, x + 0.05, 1.99, 2.95, 0.55,
        size=15, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
    txt(s, desc, x + 0.12, 2.72, 2.82, 4.1,
        size=13, color=C_DGRAY, align=PP_ALIGN.CENTER)
    if i < 3:
        txt(s, "▶", 3.33 + i * 3.2, 3.8, 0.38, 0.5,
            size=18, color=C_WARM, align=PP_ALIGN.CENTER)

rect(s, 0.45, 7.07, 12.4, 0.36, fill=C_SAND)
txt(s, "カウンセリングはクライアント自身が「なりたい自分」を言語化できるよう伴走するプロセス",
    0.65, 7.1, 12.0, 0.32, size=13, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 8 ： Have to → Want to（40代女性向け）
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "05｜Have to → Want to ── 40代女性のあるあると変化",
       bar_color=RGBColor(0x6E, 0x58, 0x48))
rect(s, 0, 1.05, 13.33, 7.5 - 1.05, fill=C_BG2)
rect(s, 0, 1.05, 13.33, 0.06, fill=C_GOLD)

rect(s, 0.4, 1.18, 12.5, 0.52, fill=C_SAND)
txt(s, "── 40代になって、こんな気持ち、ありませんか？ ──",
    0.6, 1.23, 12.1, 0.45, size=15, bold=True,
    color=C_MOCHA, align=PP_ALIGN.CENTER, italic=True)

bubbles = [
    "「また食べちゃった…\n意志が弱い私ってダメだ」",
    "「運動しなきゃとは\n思ってるんだけど…」",
    "「わかってる、でも\nできない、がずっと続いてる」",
]
for i, t in enumerate(bubbles):
    x = 0.5 + i * 4.25
    rect(s, x, 1.83, 3.9, 1.32, fill=C_WHITE, line=C_LINE, lw=1.5)
    txt(s, t, x + 0.15, 1.93, 3.6, 0.95,
        size=13, bold=True, color=C_BROWN, align=PP_ALIGN.CENTER)

txt(s, "↑　これが「Have to（〜しなければ）」思考のサイン",
    0.5, 3.28, 12.3, 0.45, size=14, italic=True,
    color=C_MGRAY, align=PP_ALIGN.CENTER)

rect(s, 0.4, 3.82, 12.5, 0.05, fill=C_WARM)

# 対比テーブル
rect(s, 0.4, 3.95, 5.85, 0.5, fill=C_RUST)
rect(s, 7.1, 3.95, 5.85, 0.5, fill=C_SAGE)
txt(s, "Have to（〜しなきゃ）", 0.5, 3.99, 5.65, 0.42,
    size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txt(s, "Want to（〜したい！）", 7.2, 3.99, 5.65, 0.42,
    size=14, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
txt(s, "→", 6.3, 4.0, 0.8, 0.42, size=18, bold=True,
    color=C_MOCHA, align=PP_ALIGN.CENTER)

rows = [
    ("「痩せなきゃいけない」",
     "「娘の結婚式に、自信を持って立ちたい」"),
    ("「食べるのを我慢しなきゃ」",
     "「体が軽い感覚を、また味わいたい」"),
    ("「運動しなきゃ」",
     "「50代も元気でいたい。孫と走り回りたい」"),
]
for j, (hv, wt) in enumerate(rows):
    y = 4.52 + j * 0.68
    bg_l = RGBColor(0xFD, 0xF0, 0xEE) if j % 2 == 0 else C_WHITE
    bg_r = RGBColor(0xF2, 0xF7, 0xEE) if j % 2 == 0 else C_BG
    rect(s, 0.4, y, 5.85, 0.63, fill=bg_l, line=C_LINE, lw=0.5)
    rect(s, 7.1, y, 5.85, 0.63, fill=bg_r, line=C_LINE, lw=0.5)
    txt(s, hv, 0.55, y + 0.1, 5.55, 0.46, size=13,
        color=C_RUST, align=PP_ALIGN.CENTER)
    txt(s, wt, 7.25, y + 0.1, 5.55, 0.46, size=13, bold=True,
        color=C_SAGE, align=PP_ALIGN.CENTER)
    txt(s, "→", 6.3, y + 0.11, 0.8, 0.44, size=18, bold=True,
        color=C_MOCHA, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 9 ： Want to を引き出す実例
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "「Want to」を引き出す ── カウンセリングの実例",
       bar_color=C_SAGE)

rect(s, 0.8, 1.18, 11.7, 0.8, fill=C_SAND, line=C_WARM, lw=1.5)
txt(s, "困っている人の「本当に叶えたいこと」を一緒に言語化するのが、カウンセリングの出発点",
    1.0, 1.23, 11.3, 0.73, size=16, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)

steps9 = [
    ("STEP 1\n表面の悩みを聴く",
     "「体重が落ちない、\nずっとリバウンドしている」\n\n→ 否定せず、まずそのまま受け取る",
     RGBColor(0xEA, 0xD5, 0xC4)),
    ("STEP 2\n本当の気持ちを掘り下げる",
     "「どうして痩せたいんですか？」\n「叶ったらどんな気持ちに？」\n\n→「娘の成人式で一緒に\n写真を撮りたい」という言葉が出てきた",
     RGBColor(0xDC, 0xCB, 0xBC)),
    ("STEP 3\nWant to が生まれる",
     "「痩せなきゃ」が\n「娘との写真のために\n体を整えたい！」に変わる\n\n→ 行動の源泉が「義務」から\n「愛情と喜び」へ",
     RGBColor(0xCC, 0xD8, 0xC0)),
]

for i, (title, desc, bg) in enumerate(steps9):
    x = 0.5 + i * 4.25
    rect(s, x, 2.1, 3.9, 4.95, fill=C_WHITE, line=C_LINE, lw=1.5)
    rect(s, x, 2.1, 3.9, 0.75, fill=bg)
    txt(s, title, x + 0.15, 2.14, 3.62, 0.68,
        size=13, bold=True, color=C_BROWN)
    txt(s, desc, x + 0.15, 2.98, 3.62, 3.9,
        size=13, color=C_DGRAY, align=PP_ALIGN.CENTER)
    if i < 2:
        txt(s, "▶", 4.27 + i * 4.25, 4.4, 0.45, 0.55,
            size=22, color=C_WARM, align=PP_ALIGN.CENTER)

rect(s, 0.5, 7.1, 12.3, 0.32, fill=C_MOCHA)
txt(s, "「痩せたい」は入口。カウンセラーはその奥にある Want to を一緒に見つける人",
    0.7, 7.13, 12.0, 0.28, size=13, bold=True,
    color=C_WHITE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 10 ： 環境・知識・伴走者の3要素
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "06｜一人ではできない ── 必要な環境と知識")

rect(s, 0.8, 1.18, 11.7, 1.0, fill=C_SAND, line=C_WARM, lw=1.5)
txt(s, "気づきは一歩目。\n「変わる」ためには、適切な環境・知識・伴走者のサポートが必要です",
    1.0, 1.22, 11.3, 0.95, size=16, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)

elems = [
    ("環境づくり",
     ["仲間・コミュニティの存在", "記録できる仕組み",
      "定期的な振り返りの場", "家族や周囲の理解"]),
    ("知識のサポート",
     ["ファスティングの正しい方法", "体の仕組みと食の関係",
      "リバウンドを防ぐ考え方", "メンタルと食欲の関連"]),
    ("伴走するカウンセラー",
     ["定期的な対話と確認", "モチベーションの維持",
      "躓いたときの軌道修正", "次の目標の設定"]),
]

for i, (title, its) in enumerate(elems):
    x = 0.5 + i * 4.25
    rect(s, x, 2.35, 3.9, 4.7, fill=C_WHITE, line=C_LINE, lw=1.5)
    rect(s, x, 2.35, 3.9, 0.78, fill=C_SAND)
    txt(s, str(i + 1), x + 0.2, 2.4, 0.6, 0.68, size=22, bold=True,
        color=C_MOCHA, align=PP_ALIGN.CENTER)
    txt(s, title, x + 0.85, 2.44, 2.85, 0.65, size=15, bold=True, color=C_BROWN)
    rect(s, x, 3.13, 3.9, 0.52, fill=C_MOCHA)
    for j, itm in enumerate(its):
        txt(s, f"◆  {itm}", x + 0.2, 3.77 + j * 0.76, 3.55, 0.65,
            size=13, color=C_DGRAY)


# ══════════════════════════════════════════════════════════
# スライド 11 ： カウンセリングは最大の教育
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "07｜カウンセリングは「最大の教育の場」",
       bar_color=RGBColor(0x6E, 0x58, 0x48))

# 左：一般的な教育
rect(s, 0.4, 1.2, 5.8, 5.8, fill=C_WHITE, line=RGBColor(0xC0, 0xA4, 0x96), lw=1.5)
rect(s, 0.4, 1.2, 5.8, 0.62, fill=RGBColor(0xC0, 0xA4, 0x96))
txt(s, "一般的な「教える」教育", 0.55, 1.25, 5.5, 0.52,
    size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
teach = [
    "知識を一方向に伝える",
    "「正解」を外から与える",
    "受け手が受動的になりやすい",
    "「わかった」だけで終わりがち",
    "行動変容につながりにくい",
]
for j, t in enumerate(teach):
    y = 2.0 + j * 0.88
    txt(s, "・", 0.55, y + 0.1, 0.4, 0.55, size=18, color=C_MGRAY)
    txt(s, t, 0.9, y + 0.1, 5.1, 0.55, size=14, color=C_DGRAY)

txt(s, "VS", 6.27, 3.8, 0.78, 0.65, size=24, bold=True,
    color=C_MGRAY, align=PP_ALIGN.CENTER)

# 右：カウンセリングの教育
rect(s, 7.1, 1.2, 5.8, 5.8, fill=C_WHITE, line=C_WARM, lw=1.5)
rect(s, 7.1, 1.2, 5.8, 0.62, fill=C_MOCHA)
txt(s, "カウンセリングによる教育", 7.25, 1.25, 5.5, 0.52,
    size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
counsel = [
    "問いかけで自ら考えさせる",
    "自分の中から「答え」が出てくる",
    "当事者意識・主体性が生まれる",
    "「気づき」は深く記憶に残る",
    "行動変容に直結する",
]
for j, t in enumerate(counsel):
    y = 2.0 + j * 0.88
    txt(s, "◆", 7.2, y + 0.1, 0.5, 0.55, size=14, color=C_GOLD)
    txt(s, t, 7.65, y + 0.1, 5.1, 0.55, size=14, bold=True, color=C_MOCHA)

rect(s, 0.4, 7.08, 12.5, 0.35, fill=RGBColor(0x6E, 0x58, 0x48))
txt(s, "「知識を教えてもらった」より「自分で気づいた」ことのほうが、人は動く",
    0.6, 7.1, 12.1, 0.3, size=13, bold=True,
    color=C_WHITE, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 12 ： 問いかけの技術
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "08｜問いかけの技術 ── 聴く力・問う力")

# 左：NG
rect(s, 0.4, 1.2, 5.8, 5.8, fill=C_WHITE, line=RGBColor(0xC0, 0xA4, 0x96), lw=1.5)
rect(s, 0.4, 1.2, 5.8, 0.62, fill=C_RUST)
txt(s, "避けたい問いかけ", 0.55, 1.25, 5.5, 0.52,
    size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
bads = [
    ("「ちゃんと食事制限しましたか？」", "→ 詰問。相手が萎縮する"),
    ("「それはこうすればいいですよ」", "→ 答えを与えすぎ。主体性が消える"),
    ("「なんでできないんですか？」", "→ 責める印象。信頼関係が壊れる"),
    ("「〇〇すべきですよね」", "→ 押しつけ。反発を生みやすい"),
]
for j, (q, note) in enumerate(bads):
    y = 1.97 + j * 1.18
    rect(s, 0.55, y, 5.5, 1.1, fill=RGBColor(0xFD, 0xF2, 0xF0),
         line=RGBColor(0xC0, 0xA4, 0x96), lw=0.8)
    txt(s, q, 0.7, y + 0.08, 5.2, 0.46,
        size=13, bold=True, color=C_RUST)
    txt(s, note, 0.7, y + 0.6, 5.2, 0.42, size=12, color=C_MGRAY)

txt(s, "VS", 6.27, 3.8, 0.78, 0.65, size=24, bold=True,
    color=C_MGRAY, align=PP_ALIGN.CENTER)

# 右：OK
rect(s, 7.1, 1.2, 5.8, 5.8, fill=C_WHITE, line=C_WARM, lw=1.5)
rect(s, 7.1, 1.2, 5.8, 0.62, fill=C_SAGE)
txt(s, "効果的な問いかけ", 7.25, 1.25, 5.5, 0.52,
    size=16, bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
goods = [
    ("「今、どんなことを感じていますか？」", "→ 感情に寄り添うオープン質問"),
    ("「うまくいったとしたら\nどんな違いがありますか？」", "→ 理想の未来を描かせる"),
    ("「それをやってみてどうでしたか？」", "→ 体験を通じた気づきを促す"),
    ("「一番大切にしたいことは何ですか？」", "→ 価値観に触れる深い問い"),
]
for j, (q, note) in enumerate(goods):
    y = 1.97 + j * 1.18
    rect(s, 7.25, y, 5.5, 1.1, fill=C_BG2, line=C_LINE, lw=0.8)
    txt(s, q, 7.4, y + 0.08, 5.2, 0.55,
        size=13, bold=True, color=C_MOCHA)
    txt(s, note, 7.4, y + 0.68, 5.2, 0.36, size=12, color=C_MGRAY)


# ══════════════════════════════════════════════════════════
# スライド 13 ： 資格取得者へのメッセージ
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.33, 7.5, fill=C_BG2)
rect(s, 0, 0, 13.33, 0.22, fill=C_BROWN)
rect(s, 0, 7.28, 13.33, 0.22, fill=C_BROWN)
rect(s, 0, 0, 0.18, 7.5, fill=C_SAND)
rect(s, 13.15, 0, 0.18, 7.5, fill=C_SAND)
rect(s, 0.18, 0.22, 0.06, 7.06, fill=C_GOLD)
rect(s, 13.09, 0.22, 0.06, 7.06, fill=C_GOLD)

txt(s, "資格を取ったあなたへ", 0.5, 1.15, 12.3, 0.95,
    size=40, bold=True, color=C_BROWN, align=PP_ALIGN.CENTER)
rect(s, 3.0, 2.18, 7.3, 0.07, fill=C_GOLD)

msgs = [
    ("「全部答えてあげなきゃ」と思わなくていい。", C_MOCHA, False),
    ("「全部伝えてあげなきゃ」と思わなくていい。", C_MOCHA, False),
    ("", C_DGRAY, False),
    ("あなたの役割は、答えを渡すことではなく、", C_BROWN, True),
    ("クライアントが自分の答えに気づく手助けをすること。", C_BROWN, True),
]
for j, (m, c, bold) in enumerate(msgs):
    txt(s, m, 0.8, 2.38 + j * 0.78, 11.7, 0.72,
        size=19 if bold else 18, bold=bold, color=c, align=PP_ALIGN.CENTER)

rect(s, 1.5, 6.28, 10.3, 0.98, fill=C_SAND, line=C_WARM, lw=2)
txt(s, "悩みに気づかせる場をつくれるカウンセラーが、\n最大の教育者になれる。",
    1.7, 6.31, 9.9, 0.92, size=18, bold=True,
    color=C_BROWN, align=PP_ALIGN.CENTER)


# ══════════════════════════════════════════════════════════
# スライド 14 ： まとめ
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
header(s, "まとめ ── カウンセリングの本質")

summary = [
    ("答えを教える場ではない",
     "自分に足りないものに気づかせる、「気づきの場」"),
    ("「なりたい自分」を見つけるプロセスを支える",
     "クライアントが自分の言葉で理想を語れるよう伴走する"),
    ("Have to → Want to への転換",
     "内側から湧く動機が、行動を長続きさせる"),
    ("一人では変われない、だから一緒に",
     "気づきの後には、環境・知識・伴走者が必要"),
    ("カウンセリングは最大の教育の場",
     "「自分で気づいた」ことは深く記憶に残り、行動変容に直結する"),
]

for i, (title, desc) in enumerate(summary):
    y = 1.25 + i * 1.2
    rect(s, 0.4, y, 12.5, 1.1, fill=C_WHITE, line=C_LINE, lw=1)
    rect(s, 0.4, y, 0.72, 1.1, fill=C_SAND)
    txt(s, str(i + 1), 0.42, y + 0.3, 0.68, 0.5, size=20, bold=True,
        color=C_MOCHA, align=PP_ALIGN.CENTER)
    txt(s, title, 1.22, y + 0.1, 11.4, 0.48,
        size=15, bold=True, color=C_BROWN)
    txt(s, desc, 1.22, y + 0.58, 11.4, 0.45, size=13, color=C_MGRAY)

rect(s, 0.4, 7.1, 12.5, 0.32, fill=C_MOCHA)
txt(s, "悩みに気づかせ、なりたい自分を見つける場 ── それがカウンセリングの真髄",
    0.6, 7.12, 12.1, 0.28, size=13, bold=True,
    color=C_WHITE, align=PP_ALIGN.CENTER)


# 保存
out = "/home/user/kiri-gohan/counseling_v3.pptx"
prs.save(out)
print(f"Saved: {out}  /  Slides: {len(prs.slides)}")
