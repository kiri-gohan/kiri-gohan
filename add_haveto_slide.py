#!/usr/bin/env python3
"""
「Have to → Want to」スライドを既存PPTXに追加
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import copy

COLOR_PRIMARY = RGBColor(0x2E, 0x7D, 0x32)
COLOR_ACCENT = RGBColor(0x81, 0xC7, 0x84)
COLOR_WARM = RGBColor(0xF9, 0xA8, 0x25)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_DARK = RGBColor(0x1A, 0x1A, 0x2E)
COLOR_LIGHT_BG = RGBColor(0xF1, 0xF8, 0xE9)
COLOR_GRAY = RGBColor(0x75, 0x75, 0x75)
COLOR_ORANGE = RGBColor(0xE6, 0x51, 0x00)
COLOR_TEAL = RGBColor(0x00, 0x89, 0x8A)
COLOR_PINK = RGBColor(0xE9, 0x1E, 0x63)
COLOR_PURPLE = RGBColor(0x6A, 0x1B, 0x9A)

prs = Presentation("/home/user/kiri-gohan/ファスティング指導士_カウンセリング講義スライド.pptx")
BLANK_LAYOUT = prs.slide_layouts[6]


def add_rect(slide, left, top, width, height, fill_color=None, line_color=None, line_width=None):
    shape = slide.shapes.add_shape(
        1,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color:
        shape.line.color.rgb = line_color
        if line_width:
            shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_text(slide, text, left, top, width, height,
             font_size=18, bold=False, color=COLOR_DARK,
             align=PP_ALIGN.LEFT, italic=False):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txBox


# ============================================================
# 新スライド A: Have to → Want to（40代女性の「あるある」から入る）
# ============================================================
slide_a = prs.slides.add_slide(BLANK_LAYOUT)

# 背景
add_rect(slide_a, 0, 0, 13.33, 7.5, fill_color=RGBColor(0xFD, 0xF6, 0xFF))
add_rect(slide_a, 0, 0, 13.33, 1.15, fill_color=COLOR_PURPLE)
add_text(slide_a, "カウンセリングが変える「動き方」── Have to から Want to へ",
         0.4, 0.18, 12.5, 0.8,
         font_size=23, bold=True, color=COLOR_WHITE)

# 40代女性あるある ヘッダー
add_rect(slide_a, 0.4, 1.25, 12.5, 0.55, fill_color=RGBColor(0xEE, 0xE0, 0xF7))
add_text(slide_a, "── 40代になって、こんな経験ありませんか？ ──",
         0.5, 1.28, 12.2, 0.48,
         font_size=16, bold=True, color=COLOR_PURPLE, align=PP_ALIGN.CENTER)

# あるある吹き出し風（3つ）
bubbles = [
    ("「また食べちゃった…\n意志が弱い私ってダメだ」", RGBColor(0xFF, 0xE0, 0xEB)),
    ("「運動しなきゃとは\n思ってるんだけど…」", RGBColor(0xFF, 0xF0, 0xCC)),
    ("「わかってるけど\nできない、がずっと続いてる」", RGBColor(0xE8, 0xF5, 0xE9)),
]
for i, (text, bg) in enumerate(bubbles):
    x = 0.5 + i * 4.25
    add_rect(slide_a, x, 1.95, 3.9, 1.5, fill_color=bg,
             line_color=COLOR_PURPLE, line_width=1)
    add_text(slide_a, text, x + 0.15, 2.05, 3.6, 1.3,
             font_size=15, bold=True, color=COLOR_DARK, align=PP_ALIGN.CENTER)

add_text(slide_a, "↑ これが「Have to（〜しなければ）」思考のサイン",
         0.5, 3.55, 12.3, 0.5,
         font_size=15, italic=True, color=COLOR_GRAY, align=PP_ALIGN.CENTER)

# 矢印・転換ゾーン
add_rect(slide_a, 0.4, 4.15, 12.5, 0.08, fill_color=RGBColor(0xCE, 0x93, 0xD8))
add_text(slide_a, "カウンセリングで「なぜそうしたいのか」を一緒に掘り下げると……",
         0.5, 4.3, 12.2, 0.5,
         font_size=15, bold=True, color=COLOR_PURPLE, align=PP_ALIGN.CENTER)

# Have to → Want to 対比
have_wants = [
    ("Have to（義務・我慢）", "Want to（内発的な動機）"),
    ("「痩せなきゃいけない」", "「娘の結婚式に自信を持って立ちたい」"),
    ("「食べるのを我慢しなきゃ」", "「体が軽くなった感覚をまた味わいたい」"),
    ("「運動しなきゃ」", "「50代も元気でいたい。孫と走り回りたい」"),
]

add_rect(slide_a, 0.4, 4.88, 5.8, 0.45, fill_color=RGBColor(0xC6, 0x28, 0x28))
add_rect(slide_a, 7.1, 4.88, 5.8, 0.45, fill_color=COLOR_PRIMARY)
add_text(slide_a, "❌  Have to（〜しなきゃ）", 0.5, 4.9, 5.6, 0.4,
         font_size=14, bold=True, color=COLOR_WHITE, align=PP_ALIGN.CENTER)
add_text(slide_a, "✅  Want to（〜したい！）", 7.2, 4.9, 5.6, 0.4,
         font_size=14, bold=True, color=COLOR_WHITE, align=PP_ALIGN.CENTER)

for j, (have, want) in enumerate(have_wants[1:]):
    y = 5.42 + j * 0.62
    bg_h = RGBColor(0xFF, 0xEB, 0xEE) if j % 2 == 0 else RGBColor(0xFF, 0xF8, 0xF8)
    bg_w = RGBColor(0xE8, 0xF5, 0xE9) if j % 2 == 0 else RGBColor(0xF1, 0xF8, 0xE9)
    add_rect(slide_a, 0.4, y, 5.8, 0.56, fill_color=bg_h)
    add_rect(slide_a, 7.1, y, 5.8, 0.56, fill_color=bg_w)
    add_text(slide_a, have, 0.55, y + 0.07, 5.5, 0.45,
             font_size=13, color=RGBColor(0xC6, 0x28, 0x28), align=PP_ALIGN.CENTER)
    add_text(slide_a, want, 7.25, y + 0.07, 5.5, 0.45,
             font_size=13, bold=True, color=COLOR_PRIMARY, align=PP_ALIGN.CENTER)
    add_text(slide_a, "→", 6.2, y + 0.1, 0.9, 0.4,
             font_size=20, bold=True, color=COLOR_GRAY, align=PP_ALIGN.CENTER)

add_rect(slide_a, 0.4, 7.3, 12.5, 0.15, fill_color=COLOR_PURPLE)


# ============================================================
# 新スライド B: カウンセリングが「Want to」を引き出す
# ============================================================
slide_b = prs.slides.add_slide(BLANK_LAYOUT)

add_rect(slide_b, 0, 0, 13.33, 7.5, fill_color=RGBColor(0xE8, 0xF5, 0xE9))
add_rect(slide_b, 0, 0, 13.33, 1.15, fill_color=COLOR_PRIMARY)
add_text(slide_b, "「Want to」を引き出すのが、カウンセラーの仕事",
         0.4, 0.18, 12.5, 0.8,
         font_size=26, bold=True, color=COLOR_WHITE)

# 核心メッセージ
add_rect(slide_b, 0.8, 1.25, 11.7, 1.05, fill_color=COLOR_PRIMARY)
add_text(slide_b,
         "「困っている人を探す」── その人が本当に叶えたいことを\n一緒に言語化することが、カウンセリングの出発点",
         1.0, 1.28, 11.3, 1.0,
         font_size=18, bold=True, color=COLOR_WHITE, align=PP_ALIGN.CENTER)

# 3ステップ（40代女性向けシナリオ）
add_text(slide_b, "実例：体重が落ちないと悩む40代女性へのカウンセリング",
         0.5, 2.45, 12.3, 0.5,
         font_size=15, bold=True, color=COLOR_TEAL)
add_rect(slide_b, 0.5, 2.5, 12.3, 0.05, fill_color=COLOR_TEAL)

steps = [
    ("STEP 1\n表面の悩みを聞く",
     "「体重が減らない、ずっとリバウンドしている」\n→ 否定せず、まずそのまま受け取る",
     RGBColor(0xFF, 0xB7, 0x4D), "😔"),
    ("STEP 2\n本当の気持ちを掘り下げる",
     "「どうして痩せたいんですか？」\n「もし叶ったら、どんな気持ちになりますか？」\n→ 「娘の成人式で一緒に写真を撮りたい」という言葉が出てきた",
     COLOR_TEAL, "💬"),
    ("STEP 3\nWant toが生まれる",
     "「痩せなきゃ」が\n「娘との写真のために、体を整えたい！」に変わる\n→ 行動の源泉が「義務」から「愛情と喜び」へ",
     COLOR_PRIMARY, "💚"),
]

for i, (title, desc, color, icon) in enumerate(steps):
    x = 0.5 + i * 4.25
    y = 3.1
    add_rect(slide_b, x, y, 3.9, 3.75, fill_color=COLOR_WHITE,
             line_color=color, line_width=2)
    add_rect(slide_b, x, y, 3.9, 0.75, fill_color=color)
    add_text(slide_b, icon, x + 0.05, y + 0.05, 0.75, 0.65,
             font_size=26, align=PP_ALIGN.CENTER)
    add_text(slide_b, title, x + 0.75, y + 0.08, 3.0, 0.65,
             font_size=12, bold=True, color=COLOR_WHITE)
    add_text(slide_b, desc, x + 0.15, y + 0.88, 3.6, 2.85,
             font_size=13, color=COLOR_DARK)
    if i < 2:
        add_text(slide_b, "→", 4.25 + i * 4.25, y + 1.85, 0.5, 0.5,
                 font_size=22, bold=True, color=COLOR_GRAY, align=PP_ALIGN.CENTER)

add_rect(slide_b, 0.5, 6.97, 12.3, 0.47, fill_color=COLOR_PRIMARY)
add_text(slide_b,
         "「痩せたい」は入口。カウンセラーはその奥にある「Want to」を一緒に見つける人",
         0.7, 6.99, 12.0, 0.44,
         font_size=14, bold=True, color=COLOR_WHITE, align=PP_ALIGN.CENTER)


output_path = "/home/user/kiri-gohan/ファスティング指導士_カウンセリング講義スライド.pptx"
prs.save(output_path)
print(f"Saved: {output_path}")
print(f"Total slides: {len(prs.slides)}")
